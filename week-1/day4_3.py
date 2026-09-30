from docx import Document
import os
from dotenv import load_dotenv
from openai import OpenAI


def extract_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    # Include text inside tables
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(parts)

text = extract_text("Agnit_Chatterjee_Resume_2026.docx")
#print(text)




load_dotenv(override=True)
api_key = os.getenv('OPENAI_API_KEY')

if not api_key:
    print("No API key was found - please head over to the troubleshooting notebook in this folder to identify & fix!")
elif not api_key.startswith("sk-proj-"):
    print("An API key was found, but it doesn't start sk-proj-; please check you're using the right key - see troubleshooting notebook")
else:
    print("API key found and looks good so far!")


openai = OpenAI()

user_content = "I want you to \
    1) Look at my profile and generate a one year plan for me in details so as to either 2x my salary or get a better job in the next 12 months. \
    2) Please consider my current skills, experience, and the latest trends in the industry. \
    3) Provide a month-by-month breakdown of actionable steps, including learning resources, networking strategies, and any certifications that would be beneficial. \
    4) Also, suggest ways to measure progress and adjust the plan as needed. My CV is as follows (in text format) \n" + text + " \
    5) Provide the output in a clear and business-like format, suitable for presentation to a mentor or career coach in .md format. \
    6) Help with deatiled links to study guides to udemy/youtube or other stydy resources and clarly articulate whihc course is best for which skill. "

messages = [
    {"role": "system", "content": "You are a professional at helping data engineers and solutions architect to grow their careers. You are an expert in career planning, skill development, and industry trends."},
    {"role": "user", "content": user_content}
    ]

response = openai.chat.completions.create(model="gpt-4.1-mini", messages=messages)

#print(response.choices[0].message.content)

with open("output.md", "w", encoding="utf-8") as file:
    file.write(response.choices[0].message.content)

print("Markdown file created successfully!")