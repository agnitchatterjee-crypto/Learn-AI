from docx import Document
import os
from dotenv import load_dotenv
from openai import OpenAI
import json
from rich.live import Live
from rich.markdown import Markdown

load_dotenv(override=True)
api_key = os.getenv('OPENAI_API_KEY')

if not api_key:
    print("No API key was found - please head over to the troubleshooting notebook in this folder to identify & fix!")
elif not api_key.startswith("sk-proj-"):
    print("An API key was found, but it doesn't start sk-proj-; please check you're using the right key - see troubleshooting notebook")
else:
    print("API key found and looks good so far!")

def extract_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]

    # Include text inside tables
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))

    return "\n".join(parts)

resume = extract_text("Agnit_Chatterjee_Resume_2026.docx")


with open('course_contents.txt', 'r', encoding='utf-8') as file:
    # Read the entire content into a string variable
    file_contents = file.read()

def professional_builder(cv_link):

    openai = OpenAI()

    text = extract_text(cv_link)

    user_content = "I want you to \
    1) Look at my profile. \
    2) Please consider my current skills, experience, and the latest trends in the industry. \
    3) I am currently working on a LLM course whose contents can be found here: " + file_contents + " \
    4) Contents of my CV: \n" + resume + " \
    5) Help come up with ideas around real world projects I can build which mirrors kind of work I am doing\
    6) Help with detailed explainataions and steps. Link to what each section of the course map to whihc section of the project to be built \
    7) If within the markdowns, you are having wide tables, switch the bullet lines instead of tables. \
    8) Please also provide a list of relevant links to study guides, courses, and resources"

    messages = [
    {"role": "system", "content": "You are a professional CV reviewer and career guide, expert at helping data engineers and solutions architect to grow their careers. You are an expert in career planning, skill development, and industry trends."},
    {"role": "user", "content": user_content}
    ]

    stream = openai.chat.completions.create(
        model="gpt-4.1-mini", 
        messages=messages,
        response_format={"type": "text"},
        stream=True
    )
    response = ""
    with Live(Markdown(""), refresh_per_second=8) as live:
        for chunk in stream:
            response += chunk.choices[0].delta.content or ""
            live.update(Markdown(response))
    return response

career_plan= professional_builder("Agnit_Chatterjee_Resume_2026.docx")

with open("output.md", "w", encoding="utf-8") as file:
    file.write(career_plan)

print("Markdown file created successfully!")